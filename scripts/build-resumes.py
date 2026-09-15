"""Build the English and Russian resume sources with python-docx.

Use the Codex bundled Python runtime. Render with the documents skill's
render_docx.py --emit_pdf, inspect both pages, then copy PDFs to public/resume.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
BASE = 'https://kidpudel.github.io/personal-website/'

CONTENT = {
    'en': {
        'name': 'IGOR KUPCHINENKO', 'role': 'PRODUCT DESIGNER',
        'summary': 'Product designer who researches, designs and ships. Software-engineering background: I prototype in code and direct AI coding agents to turn my specifications into working products. Shipped a Flutter ordering app on iOS and Android, released a macOS app, and coded an Instagram Saves concept from eight interviews.',
        'experience': 'RELEVANT DESIGN EXPERIENCE',
        'supergood': ('PizzaSushiWok (SuperGood)', 'Product Designer / Mobile Engineer | Jan 2023 - Jul 2024', [
            'Designed and built the food-ordering app, from menu and basket through payment and delivery tracking: one Flutter app replacing separate Android and iOS apps, with no design-to-code handoff.',
            'Customer conversations and ordering tasks highlighted uncertainty about dish details, delivery time and total cost. Added fuller dish information and a persistent order summary.',
            'Made previous orders reusable: users could add individual dishes or the whole order back to the basket. Added access through favourites and adapted layouts for different mobile screens.',
            'Built the app theme on Material Design with custom components bound to it, including a menu whose category tabs and dish list scroll in sync, with animated transitions.',
        ]),
        'projects_title': 'SELECTED PROJECTS',
        'projects': [
            ('Instagram Saves redesign', 'Independent concept, coded prototype | 2026', 'instagram-saves-redesign', [
                'Interviewed eight people about saving and finding posts. Turned the recurring problems into collection suggestions, search and date filters, notes, pins and shared reactions.',
                'Built the prototype in code (React): an interactive app running live on my site, plus walkthrough films in Remotion. Explored surfacing saved collections in the main feed.',
            ]),
            ('Observatory', 'Personal project, released on GitHub | Jul 2026 - Aug 2026', 'observatory', [
                'Five exploratory interviews shaped an app-first view of Mac performance for developers and power users: processes grouped into application totals, with detail one level below.',
                'Designed the interface and saved-recording comparisons, wrote the specifications and directed AI coding agents through the SwiftUI build. Released 0.1.1 with identity, artwork and launch film.',
            ]),
            ('Two Sticks', 'Collaborative prototype | May 2024 - Jun 2024', 'two-sticks', [
                "Turned a diploma team's pedagogical research into a Telegram flow for finding, saving and practising Chinese characters. Designed and built the bot and handwriting web app; no longer hosted.",
            ]),
        ],
        'additional_title': 'ADDITIONAL EXPERIENCE',
        'additional': [
            ('Paycos', 'Software Developer, contract | Jul 2024 - Dec 2025', 'Built payment and order-processing services. Cut PDF receipt processing from 6-10 s to 650 ms by extracting text directly, keeping OCR for images.'),
            ('22bytes', 'Mobile Game Designer / Prototype Developer | Nov 2022 - Jan 2023', 'Built playable Android prototypes in short cycles, shaping interaction flows, player feedback and scope.'),
        ],
        'skills_title': 'SKILLS',
        'skills': 'Design: user interviews, synthesis, flows, interaction and visual design, motion, prototyping, usability testing. Figma, Material Design.\nCode: HTML, CSS, JavaScript, React and Flutter prototypes; a SwiftUI app shipped by directing AI coding agents (Claude Code, Codex); Go, Python; Unity, Godot, raylib, OpenGL.\nPortfolio site built by me in Astro and React, including the embedded interactive Instagram prototype.',
        'education_title': 'EDUCATION AND LANGUAGES',
        'education': 'Bachelor of Computer & Information Science, Moscow Finance and Law Academy (MFUA) | 2019 - 2023',
        'languages': 'Russian: native | English: B2 (upper-intermediate)',
    },
    'ru': {
        'name': 'ИГОРЬ КУПЧИНЕНКО', 'role': 'ПРОДУКТОВЫЙ ДИЗАЙНЕР',
        'summary': 'Продуктовый дизайнер: исследую, проектирую и довожу до релиза. Опыт разработки: собираю прототипы в коде и направляю AI-агентов, превращая свои спецификации в работающий продукт. Выпустил Flutter-приложение для заказа еды на iOS и Android и приложение для macOS, собрал в коде концепт сохранённых Instagram по восьми интервью.',
        'experience': 'ОПЫТ В ПРОДУКТОВОМ ДИЗАЙНЕ',
        'supergood': ('PizzaSushiWok (SuperGood)', 'Продуктовый дизайнер / мобильный разработчик | январь 2023 - июль 2024', [
            'Спроектировал и реализовал заказ от меню до оплаты и отслеживания доставки: одно Flutter-приложение вместо отдельных iOS и Android, без передачи макетов в разработку.',
            'Беседы с пользователями и проверки сценария заказа выявили нехватку информации о блюдах, доставке и стоимости. Добавил подробные карточки блюд и закреплённую сводку заказа.',
            'Сделал повторный заказ из истории: можно добавить отдельные блюда или весь заказ в корзину. Добавил переход из избранного и адаптировал интерфейс под разные мобильные экраны.',
            'Собрал тему приложения на Material Design и привязанные к ней кастомные компоненты, включая меню с синхронной прокруткой категорий и блюд и анимациями.',
        ]),
        'projects_title': 'ИЗБРАННЫЕ ПРОЕКТЫ',
        'projects': [
            ('Поиск сохранённых публикаций', 'Самостоятельный концепт, прототип в коде | 2026', 'instagram-saves-redesign', [
                'Провёл восемь интервью о том, как люди сохраняют и находят посты. Превратил повторяющиеся проблемы в подсказки коллекций, поиск, фильтры по дате, заметки, закрепление и реакции.',
                'Собрал прототип в коде (React): интерактивное приложение прямо на моём сайте и ролики-прохождения в Remotion. Продумал возврат сохранённых коллекций в основную ленту.',
            ]),
            ('Observatory', 'Личный проект, выпущен на GitHub | июль 2026 - август 2026', 'observatory', [
                'Пять интервью привели к просмотру нагрузки Mac по приложениям для разработчиков и опытных пользователей: процессы сгруппированы в итог по приложению, детали уровнем ниже.',
                'Спроектировал интерфейс и сравнение записей, написал спецификации и направлял AI-агентов в реализации на SwiftUI. Выпустил 0.1.1 с айдентикой, иллюстрацией и релизным роликом.',
            ]),
            ('Two Sticks', 'Командный прототип | май 2024 - июнь 2024', 'two-sticks', [
                'На основе исследования дипломной команды МПГУ спроектировал поиск, сохранение и практику иероглифов в Telegram. Реализовал бота и веб-приложение для письма; уже не размещён.',
            ]),
        ],
        'additional_title': 'ДРУГОЙ ОПЫТ РАБОТЫ',
        'additional': [
            ('Paycos', 'Разработчик ПО, контракт | июль 2024 - декабрь 2025', 'Разрабатывал сервисы оплаты и обработки заказов. Сократил обработку PDF-чеков с 6-10 с до 650 мс: извлекал текст напрямую, оставив OCR для изображений.'),
            ('22bytes', 'Гейм-дизайнер / разработчик прототипов | ноябрь 2022 - январь 2023', 'Быстро собирал игровые Android-прототипы: взаимодействие, фидбек игроку, объём работы.'),
        ],
        'skills_title': 'НАВЫКИ',
        'skills': 'Дизайн: интервью, сценарии, UX/UI, анимация, прототипы, юзабилити-тесты. Figma, Material Design.\nКод: прототипы на HTML, CSS, JavaScript, React, Flutter; SwiftUI-продукт, собранный с AI-агентами (Claude Code, Codex); Go, Python; Unity, Godot, raylib, OpenGL.\nПортфолио-сайт собрал сам на Astro и React, включая интерактивный прототип Instagram.',
        'education_title': 'ОБРАЗОВАНИЕ И ЯЗЫКИ',
        'education': 'Бакалавр компьютерных и информационных наук | МФЮА | 2019 - 2023',
        'languages': 'Русский: родной | Английский: B2 (Upper-Intermediate)',
    },
}


def link(p, text, url, size=10.5, bold=False):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    for tag, val in [('w:rFonts', None), ('w:color', '222222'), ('w:sz', str(int(size*2))), ('w:u', 'single')]:
        x = OxmlElement(tag)
        if val is not None: x.set(qn('w:val'), val)
        else:
            for attr in ['ascii','hAnsi','cs']: x.set(qn('w:'+attr), 'Arial')
        pr.append(x)
    if bold: pr.append(OxmlElement('w:b'))
    r.append(pr); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)


def build(lang):
    c=CONTENT[lang]; d=Document(); sec=d.sections[0]
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=Mm(11); sec.bottom_margin=Mm(11)
    sec.left_margin=Mm(16); sec.right_margin=Mm(16)
    normal=d.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10.5)
    normal.font.color.rgb=RGBColor.from_string('222222')
    normal.paragraph_format.line_spacing=1.04
    normal.paragraph_format.space_after=Pt(3)
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ['Title','Subtitle','Heading 1','Heading 2','List Bullet']:
        s=d.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    d.styles['Title'].font.size=Pt(23); d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(1)
    d.styles['Subtitle'].font.size=Pt(11); d.styles['Subtitle'].font.bold=True
    d.styles['Subtitle'].font.italic=False
    d.styles['Subtitle'].paragraph_format.space_after=Pt(5)
    h=d.styles['Heading 1']; h.font.size=Pt(10.5); h.font.bold=True
    h.paragraph_format.space_before=Pt(4); h.paragraph_format.space_after=Pt(2)
    h.paragraph_format.keep_with_next=True
    h2=d.styles['Heading 2']; h2.font.size=Pt(10.5); h2.font.bold=True
    h2.paragraph_format.space_before=Pt(4); h2.paragraph_format.space_after=Pt(1)
    h2.paragraph_format.keep_with_next=True
    bullet=d.styles['List Bullet']; bullet.font.size=Pt(10.5)
    bullet.paragraph_format.left_indent=Mm(3); bullet.paragraph_format.first_line_indent=Mm(-3)
    bullet.paragraph_format.space_after=Pt(2)
    d.core_properties.author=c['name']; d.core_properties.title=c['name']+' '+c['role']
    d.core_properties.language='ru-RU' if lang=='ru' else 'en-US'
    settings=d.settings.element; default=settings.find(qn('w:themeFontLang'))
    if default is not None: default.set(qn('w:val'), 'ru-RU' if lang=='ru' else 'en-US')
    d.add_paragraph(c['name'],'Title'); d.add_paragraph(c['role'],'Subtitle')
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    link(p,'i.kupchinenko@gmail.com','mailto:i.kupchinenko@gmail.com',10)
    p.add_run('  |  ')
    link(p,'kidpudel.github.io/personal-website',BASE,10)
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(6)
    link(p,'linkedin.com/in/iggydev','https://www.linkedin.com/in/iggydev',10)
    p.add_run('  |  '); link(p,'github.com/KidPudel','https://github.com/KidPudel',10)
    d.add_paragraph(c['summary'])

    def entry(title,context,body,slug=None):
        p=d.add_paragraph(style='Heading 2')
        if slug: link(p,title,BASE+'case-studies/'+slug+'/',bold=True)
        else: p.add_run(title)
        p=d.add_paragraph(context); p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(2)
        for r in p.runs: r.font.size=Pt(9.5)
        for s in body:
            p=d.add_paragraph(s,'List Bullet'); p.paragraph_format.keep_together=True

    d.add_paragraph(c['experience'],'Heading 1')
    title,ctx,bullets=c['supergood']; entry(title,ctx,bullets,'supergood')
    d.add_paragraph(c['projects_title'],'Heading 1')
    for title,ctx,slug,bullets in c['projects']: entry(title,ctx,bullets,slug)
    d.add_paragraph(c['additional_title'],'Heading 1')
    for title,ctx,body in c['additional']: entry(title,ctx,[body])
    d.add_paragraph(c['skills_title'],'Heading 1'); d.add_paragraph(c['skills'])
    d.add_paragraph(c['education_title'],'Heading 1')
    p=d.add_paragraph(c['education']); p.paragraph_format.space_after=Pt(1)
    d.add_paragraph(c['languages'])
    name='igor-kupchinenko-product-designer-resume-2026.docx' if lang=='en' else 'igor-kupchinenko-product-designer-resume-ru.docx'
    target=ROOT/'public'/'resume'/name; d.save(target); print(target)


if __name__ == '__main__':
    for lang in CONTENT: build(lang)
